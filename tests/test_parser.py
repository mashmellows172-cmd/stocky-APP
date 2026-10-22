from pathlib import Path

import pytest

from stocky.parser import ParseError, parse_rows, read_report

FIXTURE = Path(__file__).parent / "fixtures" / "reporte_recorte.txt"

ITEM_26726760 = (
    " |  1   |7881 MEDIANA                |26726760  |00200026726766 |"
    " AMERICAN TOURISTER MALETA FRANKFURT. |(UNI)   -1|(UNI)   -1|       -1 |       -1 |"
    "     59990 |     -49994|"
)
ITEM_26726760_CON_CODIGO_2 = (
    " |  1   |7881 MEDIANA                |26726760  |00200026726799 |"
    " AMERICAN TOURISTER MALETA FRANKFURT. |(UNI)   -1|(UNI)   -1|       -1 |       -1 |"
    "     59990 |     -49994|"
)
ITEM_26726760_STOCK_5 = (
    " |  1   |7881 MEDIANA                |26726760  |00200026726766 |"
    " AMERICAN TOURISTER MALETA FRANKFURT. |(UNI)    5|(UNI)    5|        5 |        5 |"
    "     59990 |     -49994|"
)
CONT_07501068898989 = (
    " |      |                            |          |07501068898989 |"
    "                                      |          |          |          |          |           |           |"
)


def _fixture_lines() -> list[str]:
    return FIXTURE.read_text(encoding="utf-8").splitlines()


def test_parsea_un_item_completo():
    report = parse_rows(_fixture_lines())
    row = report.rows[0]
    assert row.area == 1
    assert row.rubro == "7881 MEDIANA"
    assert row.plu == "26726760"
    assert row.barcode == "00200026726766"
    assert row.desc == "AMERICAN TOURISTER MALETA FRANKFURT."
    assert (row.existencia, row.disponible, row.precio, row.costo) == (-1, -1, 59990, -49994)


def test_ignora_encabezados_y_reglas():
    report = parse_rows(_fixture_lines())
    assert report.unknown_lines == 0
    assert len(report.rows) == 3


def test_conserva_ceros_del_codigo():
    report = parse_rows(_fixture_lines())
    assert report.rows[0].barcode.startswith("0020")


def test_asocia_codigo_alternativo_al_item_anterior():
    report = parse_rows(_fixture_lines())
    assert report.rows[0].alt_barcodes == ["07501068898989"]
    assert report.rows[1].alt_barcodes == []


def test_lee_metadatos_del_encabezado():
    report = parse_rows(_fixture_lines())
    assert report.dependencia == "10 MAIPU"
    assert report.fecha == "17/06/2026"
    assert report.hora == "10:06:36"


def test_continuacion_sin_item_previo_cuenta_como_desconocida():
    report = parse_rows([CONT_07501068898989, ITEM_26726760])
    assert report.unknown_lines == 1
    assert report.rows[0].alt_barcodes == []


def test_salta_encabezados_de_pagina_intermedios():
    report = parse_rows(_fixture_lines())
    assert len(report.rows) == 3
    assert report.unknown_lines == 0

REAL = Path(r"C:\Users\victus\Downloads\stock17-06.xlsx")


@pytest.mark.slow
@pytest.mark.skipif(not REAL.exists(), reason="archivo real no disponible")
def test_reporte_real_completo():
    report = read_report(REAL)
    assert len(report.rows) == 26460  # 26.527 líneas - 67 PLU fusionados
    assert report.conflicts == 0
    assert report.unknown_lines == 0
    assert len({row.area for row in report.rows}) == 69
    assert len(report.areas()) == 68  # el área 50 no tiene stock disponible
    assert sum(len(r.alt_barcodes) for r in report.rows) == 3259  # 3.192 + 67 fusionados
    assert any(r.plu == "23842742" for r in report.rows)  # fila desbordada a B
    assert (report.dependencia, report.fecha, report.hora) == ("10 MAIPU", "17/06/2026", "10:06:36")


def test_reporte_vacio_lanza_error():
    with pytest.raises(ParseError, match="no trae ítems"):
        parse_rows([",".join(["x"] * 13)])


def test_plu_repetido_con_otro_codigo_se_fusiona():
    report = parse_rows([ITEM_26726760, CONT_07501068898989, ITEM_26726760_CON_CODIGO_2])
    assert len(report.rows) == 1
    assert report.rows[0].barcode == "00200026726766"
    assert "00200026726799" in report.rows[0].alt_barcodes
    assert report.conflicts == 0


def test_plu_repetido_con_stock_distinto_es_conflicto():
    report = parse_rows([ITEM_26726760, ITEM_26726760_STOCK_5])
    assert len(report.rows) == 1
    assert report.rows[0].disponible == -1
    assert report.conflicts == 1


def test_reporte_de_otro_formato_lanza_error(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active["A1"] = "AREA"
    wb.active["B1"] = "PLU"
    p = tmp_path / "otro.xlsx"
    wb.save(p)
    with pytest.raises(ParseError, match="Formato no soportado"):
        read_report(p)


def test_linea_con_campos_inesperados_cuenta_como_desconocida():
    report = parse_rows([ITEM_26726760, ITEM_26726760 + "|"])
    assert len(report.rows) == 1
    assert report.unknown_lines == 1


def _write_column_report(path: Path) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    filas = [
        [None, None, None, "FECHA : 19/10/2026"],
        [None, None, None, "HORA  : 10:14:33"],
        [None, None, None, "DEPENDENCIA   :   10 MAIPU"],
        [None, "AREA", "NIVEL DE AGREGACION", "PLU", "CODIGO DE BARRA", "DESCRIPCION PRODUCTO", "DISPONIBLE", "PRECIO DE"],
        [None, 9, "7565 TELEFONIA", "26726760", "00200026726766", "CELULAR ALFA", 2, 59990],
        [None, 9, "7565 TELEFONIA", "26726760", "07501068898989", "CELULAR ALFA", 2, 59990],
        [None, None, None, None, "07501068898989", None, None, None],
        [None, "AREA", "NIVEL DE AGREGACION", "PLU", "CODIGO DE BARRA", "DESCRIPCION PRODUCTO", "DISPONIBLE", "PRECIO DE"],
        [None, 9, "7565 TELEFONIA", "26726761", "00200026726799", "CELULAR BETA", 0, 37990],
        [None, 8, "7881 AUDIO", "26726762", "00200026726762", "PARLANTE | GAMMA", 1, 15000],
        [None, 10, "9999 VARIOS", "26726764", "00200026726764", "REPUESTO", None, None],
    ]
    for fila in filas:
        ws.append(fila)
    wb.save(path)
    return path


def test_lee_reporte_en_columnas(tmp_path):
    p = _write_column_report(tmp_path / "columnas.xlsx")
    report = read_report(p)
    assert (report.dependencia, report.fecha, report.hora) == ("10 MAIPU", "19/10/2026", "10:14:33")
    assert len(report.rows) == 4
    assert report.unknown_lines == 0
    assert report.conflicts == 0
    alfa = report.rows[0]
    assert (alfa.area, alfa.rubro, alfa.plu, alfa.barcode, alfa.desc) == (
        9,
        "7565 TELEFONIA",
        "26726760",
        "00200026726766",
        "CELULAR ALFA",
    )
    assert (alfa.un_emp, alfa.un_despa, alfa.existencia, alfa.disponible, alfa.precio, alfa.costo) == (0, 0, 0, 2, 59990, 0)


def test_columnas_ignoran_encabezados_y_ruido(tmp_path):
    p = _write_column_report(tmp_path / "columnas.xlsx")
    report = read_report(p)
    assert len(report.rows) == 4
    assert report.unknown_lines == 0
    assert [fila.desc for fila in report.rows] == ["CELULAR ALFA", "CELULAR BETA", "PARLANTE | GAMMA", "REPUESTO"]


def test_columnas_fusionan_plu_repetido(tmp_path):
    p = _write_column_report(tmp_path / "columnas.xlsx")
    report = read_report(p)
    alfa = report.rows[0]
    assert alfa.plu == "26726760"
    assert alfa.alt_barcodes == ["07501068898989"]
    assert len([fila for fila in report.rows if fila.plu == "26726760"]) == 1


def test_columnas_con_campos_vacios_quedan_en_cero(tmp_path):
    p = _write_column_report(tmp_path / "columnas.xlsx")
    report = read_report(p)
    repuesto = report.rows[3]
    assert (repuesto.disponible, repuesto.precio) == (0, 0)
    assert report.areas() == {8: 1, 9: 1}


def test_columnas_con_pipe_en_descripcion_no_se_confunde(tmp_path):
    p = _write_column_report(tmp_path / "columnas.xlsx")
    report = read_report(p)
    assert report.rows[2].desc == "PARLANTE | GAMMA"


def test_tabla_sin_encabezado_no_se_adivina(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.append([None, 9, "7565 TELEFONIA", "26726760", "00200026726766", "CELULAR ALFA", 2, 59990])
    p = tmp_path / "sin_encabezado.xlsx"
    wb.save(p)
    with pytest.raises(ParseError, match="Formato no soportado"):
        read_report(p)


REAL_OCTUBRE = Path(r"C:\Users\victus\Downloads\stock19-10.xlsx")


@pytest.mark.slow
@pytest.mark.skipif(not REAL_OCTUBRE.exists(), reason="archivo real no disponible")
def test_reporte_de_columnas_real():
    report = read_report(REAL_OCTUBRE)
    assert len(report.rows) == 30424
    assert report.conflicts == 0
    assert report.unknown_lines == 0
    assert len({row.area for row in report.rows}) == 66
    assert len(report.areas()) == 66
    assert report.areas()[8] == 39
    assert report.areas()[9] == 107
    assert report.areas()[29] == 133
    assert report.areas()[46] == 36
    assert report.areas()[84] == 17
    assert report.areas()[99] == 74
    assert sum(row.disponible for row in report.rows) == 299342
    assert (report.dependencia, report.fecha, report.hora) == ("10 MAIPU", "19/10/2026", "10:14:33")
    primero = report.rows[0]
    assert (primero.area, primero.plu, primero.desc, primero.disponible, primero.precio) == (
        98,
        "25985176",
        "BABY MINK BM322 BEBE.RAYAS 9/12MESES",
        1,
        4450,
    )
