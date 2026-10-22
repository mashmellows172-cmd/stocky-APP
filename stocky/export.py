"""Exportación a xlsx lista para imprimir."""

from __future__ import annotations

import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from .areas import file_stem, header_label
from .models import StockReport, StockRow
from .presets import Formato

TITLE_ROW = 1
META_ROW = 2
HEADER_ROW = 4
FIRST_DATA_ROW = 5

ROW_HEIGHTS = {"compacta": 12, "normal": 15, "amplia": 18}
HEADER_COLORS = {"gris": "D9D9D9", "verde": "C6EFCE"}
ZEBRA_COLOR = "F2F2F2"
THIN = Side(style="thin", color="808080")


def _sort_key(row: StockRow) -> tuple[str, str]:
    plain = unicodedata.normalize("NFKD", row.desc).encode("ascii", "ignore").decode().upper()
    return plain, row.plu


def rows_for(report: StockReport, area: int, min_stock: int = 1) -> list[StockRow]:
    """Ítems del área con stock suficiente, ordenados por descripción."""
    selected = [row for row in report.rows if row.area == area and row.disponible >= min_stock]
    return sorted(selected, key=_sort_key)


def _number_format(fmt: Formato) -> str:
    return "#,##0" if fmt.thousands else "0"


def _header_fill(fmt: Formato):
    color = HEADER_COLORS.get(fmt.header_fill)
    return PatternFill("solid", fgColor=color) if color else None


def _text(ws, row: int, column: int, value: str, font: Font, align: str = "left") -> None:
    cell = ws.cell(row=row, column=column, value=value)
    cell.number_format = "@"
    cell.font = font
    cell.alignment = Alignment(horizontal=align, vertical="center")


def build_workbook(rows: list[StockRow], area: int, report: StockReport, fmt: Formato) -> Workbook:
    workbook = Workbook()
    ws = workbook.active
    ws.title = f"AREA {area:02d}"[:31]

    headers = ["PLU", "CÓDIGO"]
    widths = [11, 17]
    if fmt.show_rubro:
        headers.append("RUBRO")
        widths.append(24)
    headers += ["DESCRIPCIÓN", "DISPONIBLE", "PRECIO"]
    widths += [56 if fmt.wide_desc else 44, 11, 13]
    last_column = len(headers)

    # Las columnas se miden en caracteres del tamaño base (11 pt en Excel); al
    # agrandar la letra hay que ensancharlas o el texto se recorta al ajustar la
    # hoja a una página de ancho.
    factor = fmt.font_size / 11
    widths = [round(width * factor, 1) for width in widths]

    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    title_font = Font(name=fmt.font_family, size=fmt.title_size, bold=True)
    _text(ws, TITLE_ROW, 1, header_label(area), title_font)
    if fmt.extra_header:
        _text(ws, TITLE_ROW, 1, f"{header_label(area)} · {fmt.extra_header}", title_font)
    ws.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=last_column)

    units = sum(row.disponible for row in rows)
    dependencia = report.dependencia or "SIN DEPENDENCIA"
    meta = " · ".join(
        part for part in (dependencia, report.fecha, report.hora, f"{len(rows)} ítems", f"{units} unidades") if part
    )
    meta_font = Font(name=fmt.font_family, size=max(8, fmt.font_size - 1))
    _text(ws, META_ROW, 1, meta, meta_font)
    ws.merge_cells(start_row=META_ROW, start_column=1, end_row=META_ROW, end_column=last_column)

    header_font = Font(name=fmt.font_family, size=fmt.font_size, bold=True)
    fill = _header_fill(fmt)
    border = Border(bottom=THIN)
    for index, header in enumerate(headers, start=1):
        cell = ws.cell(row=HEADER_ROW, column=index, value=header)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        if fill is not None:
            cell.fill = fill

    body_font = Font(name=fmt.font_family, size=fmt.font_size)
    alt_font = Font(name=fmt.font_family, size=max(6, fmt.font_size - 2))
    number_format = _number_format(fmt)
    height = ROW_HEIGHTS.get(fmt.density, 15)
    zebra = PatternFill("solid", fgColor=ZEBRA_COLOR)

    current = FIRST_DATA_ROW
    for position, row in enumerate(rows):
        cells = {
            1: (row.plu, "center"),
            2: (row.barcode, "center"),
        }
        following = 3
        if fmt.show_rubro:
            cells[following] = (row.rubro, "left")
            following += 1
        cells[following] = (row.desc, "left")
        cells[following + 1] = (row.disponible, "center")
        cells[following + 2] = (row.precio, "center")

        for column, (value, align) in cells.items():
            cell = ws.cell(row=current, column=column, value=value)
            cell.font = body_font
            cell.alignment = Alignment(horizontal=align, vertical="center")
            if column in (1, 2):
                cell.number_format = "@"
            elif column >= following + 1:
                cell.number_format = number_format
            else:
                cell.number_format = "@"
            if fmt.zebra and position % 2 == 1:
                cell.fill = zebra
        ws.row_dimensions[current].height = height
        current += 1

        for alt in row.alt_barcodes:
            _text(ws, current, 2, alt, alt_font, "center")
            ws.row_dimensions[current].height = max(10, height - 3)
            current += 1

    last_row = max(current - 1, HEADER_ROW)

    ws.print_area = f"A1:{get_column_letter(last_column)}{last_row}"
    ws.print_title_rows = f"{TITLE_ROW}:{HEADER_ROW}"
    ws.sheet_view.showGridLines = False

    setup = ws.page_setup
    setup.paperSize = ws.PAPERSIZE_LETTER
    setup.orientation = "landscape" if fmt.orientation == "horizontal" else "portrait"
    setup.fitToWidth = 1
    setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    margin = fmt.margin_cm / 2.54
    ws.page_margins.left = margin
    ws.page_margins.right = margin
    ws.page_margins.top = margin
    ws.page_margins.bottom = margin

    ws.oddFooter.left.text = f"{len(rows)} ítems · {units} unidades"
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Página &P de &N"
    ws.oddFooter.right.size = 8

    return workbook


def write_area_xlsx(report: StockReport, area: int, fmt: Formato, out_dir: Path) -> Path | None:
    rows = rows_for(report, area, fmt.min_stock)
    if not rows:
        return None
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    workbook = build_workbook(rows, area, report, fmt)
    path = out_dir / f"{file_stem(area)}.xlsx"
    workbook.save(path)
    return path


def _iso_date(fecha: str) -> str:
    parts = fecha.split("/")
    if len(parts) == 3 and all(part.isdigit() for part in parts):
        day, month, year = parts
        return f"{year}-{month}-{day}"
    return "sin-fecha"


def output_dir(report: StockReport, source: Path) -> Path:
    """Carpeta de salida junto al reporte: ``Stock 2026-06-17``."""
    return Path(source).parent / f"Stock {_iso_date(report.fecha)}"
